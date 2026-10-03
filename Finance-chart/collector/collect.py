#!/usr/bin/env python3
"""금융사 연표 차트 - 데이터 수집기.

각 지표를 공식/공개 출처에서 내려받아 data/<id>.json 으로 저장하고
data/manifest.json 에 지표 목록과 수집 상태를 기록한다.

- 하루 한 번 GitHub Actions 가 실행한다 (.github/workflows/update.yml).
- 지표 하나가 실패해도 나머지는 계속 진행하며, 실패한 지표는 이전 데이터를 유지한다.
- 필요한 비밀 키(환경변수): FRED_API_KEY(없어도 동작), ECOS_API_KEY(한국은행 금리에 필요)
"""
import datetime as dt
import io
import json
import math
import os
import re
import sys
import tempfile
import time
import traceback

import requests

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DATA_DIR = os.environ.get("FC_DATA_DIR", os.path.join(ROOT, "data"))
EPOCH = dt.date(1970, 1, 1)
MIN_DATE = dt.date(1600, 1, 1)
TODAY = dt.datetime.now(dt.timezone.utc).date()
END = TODAY - dt.timedelta(days=1)  # 하루가 완결된 마지막 날까지만 저장
UA = "Mozilla/5.0 (compatible; finance-chart-collector/1.0)"
SECRETS = [os.environ.get(k) for k in ("FRED_API_KEY", "ECOS_API_KEY")]


class SkipSource(Exception):
    """키가 없는 등, 정상적으로 건너뛰는 경우."""


def scrub(msg):
    msg = str(msg)
    for s in SECRETS:
        if s:
            msg = msg.replace(s, "***")
    return msg


def log(*a):
    print(scrub(" ".join(str(x) for x in a)), flush=True)


def http_get(url, params=None, headers=None, retries=2, timeout=(10, 45), stream=False):
    h = {"User-Agent": UA}
    if headers:
        h.update(headers)
    last = None
    for i in range(retries):
        try:
            r = requests.get(url, params=params, headers=h, timeout=timeout, stream=stream)
            if r.status_code in (429, 500, 502, 503, 504):
                raise requests.HTTPError(f"HTTP {r.status_code}")
            r.raise_for_status()
            return r
        except Exception as e:  # noqa: BLE001
            last = e
            # 4xx(429 제외)는 재시도해도 소용없다
            resp = getattr(e, "response", None)
            if resp is not None and 400 <= resp.status_code < 500 and resp.status_code != 429:
                break
            time.sleep(2 * (i + 1))
    raise RuntimeError(scrub(f"요청 실패: {last}"))


def day_num(d):
    return (d - EPOCH).days


def parse_date(s):
    s = str(s).strip()
    m = re.match(r"^(\d{4})-(\d{2})-(\d{2})", s)
    if m:
        return dt.date(int(m[1]), int(m[2]), int(m[3]))
    m = re.match(r"^(\d{4})-(\d{2})$", s)
    if m:
        return dt.date(int(m[1]), int(m[2]), 1)
    m = re.match(r"^(\d{4})(\d{2})(\d{2})$", s)
    if m:
        return dt.date(int(m[1]), int(m[2]), int(m[3]))
    m = re.match(r"^(\d{4})(\d{2})$", s)
    if m:
        return dt.date(int(m[1]), int(m[2]), 1)
    m = re.match(r"^(\d{4})$", s)
    if m:
        return dt.date(int(m[1]), 7, 1)  # 연 단위 자료는 해당 연도 중간에 표시
    raise ValueError(f"날짜 형식을 알 수 없음: {s}")


def clean(points):
    """[(date, value)] -> 날짜순, 중복 제거, 범위 필터."""
    d = {}
    for dte, v in points:
        try:
            v = float(v)
        except (TypeError, ValueError):
            continue
        if not math.isfinite(v):
            continue
        if dte < MIN_DATE or dte > END:
            continue
        d[dte] = v
    return sorted(d.items())


# ---------------------------------------------------------------- 출처별 수집 함수
_fred_fail = [0]


def fred(series_id, invert=False):
    key = os.environ.get("FRED_API_KEY")
    pts = []
    if not key and _fred_fail[0] >= 2:
        raise SkipSource("FRED 에 연속으로 접속하지 못해 건너뜀 (FRED_API_KEY 등록을 권장)")
    try:
        pts = _fred_fetch(series_id, key)
    except Exception:
        _fred_fail[0] += 1
        raise
    _fred_fail[0] = 0
    return _fred_convert(pts, invert)


def _fred_fetch(series_id, key):
    pts = []
    if key:
        r = http_get(
            "https://api.stlouisfed.org/fred/series/observations",
            params={"series_id": series_id, "api_key": key, "file_type": "json"},
        )
        for o in r.json()["observations"]:
            pts.append((parse_date(o["date"]), o["value"]))
    else:
        r = http_get("https://fred.stlouisfed.org/graph/fredgraph.csv", params={"id": series_id})
        lines = r.text.strip().splitlines()
        for ln in lines[1:]:
            a = ln.split(",")
            if len(a) >= 2:
                pts.append((parse_date(a[0]), a[1]))
    return pts


def _fred_convert(pts, invert):
    out = []
    for d, v in pts:
        try:
            f = float(v)
        except ValueError:
            continue  # "." = 결측
        if invert:
            if f == 0:
                continue
            f = 1.0 / f
        out.append((d, f))
    return out


def ecos_base_rate():
    key = os.environ.get("ECOS_API_KEY")
    if not key:
        raise SkipSource("ECOS_API_KEY 가 설정되지 않음")
    stat, item = "722Y001", "0101000"
    for cycle, fmt_start, fmt_end in (("D", "19990101", END.strftime("%Y%m%d")), ("M", "199905", END.strftime("%Y%m"))):
        rows, start, page = [], 1, 10000
        try:
            while True:
                url = (
                    f"https://ecos.bok.or.kr/api/StatisticSearch/{key}/json/kr/"
                    f"{start}/{start + page - 1}/{stat}/{cycle}/{fmt_start}/{fmt_end}/{item}"
                )
                j = http_get(url).json()
                if "StatisticSearch" not in j:
                    raise RuntimeError(f"ECOS 응답 오류: {j.get('RESULT', j)}")
                blk = j["StatisticSearch"]
                rows += blk["row"]
                start += page
                if start > int(blk["list_total_count"]):
                    break
        except Exception as e:  # noqa: BLE001
            log(f"  ECOS 주기 {cycle} 실패: {scrub(e)}")
            continue
        if rows:
            return [(parse_date(r["TIME"]), r["DATA_VALUE"]) for r in rows]
    raise RuntimeError("ECOS 에서 기준금리를 가져오지 못함")


def binance(symbol):
    bases = ["https://data-api.binance.vision", "https://api.binance.com", "https://api1.binance.com"]
    day_ms = 86400000
    for base in bases:
        try:
            pts, start = [], 1500000000000  # 2017-07 (상장 이전부터)
            while True:
                r = http_get(
                    base + "/api/v3/klines",
                    params={"symbol": symbol, "interval": "1d", "startTime": start, "limit": 1000},
                    retries=2,
                )
                rows = r.json()
                if not rows:
                    break
                for k in rows:
                    d = dt.datetime.fromtimestamp(k[0] / 1000, dt.timezone.utc).date()
                    pts.append((d, k[4]))
                if len(rows) < 1000:
                    break
                start = rows[-1][0] + day_ms
                time.sleep(0.2)
            if pts:
                return pts
        except Exception as e:  # noqa: BLE001
            log(f"  {base} 실패: {scrub(e)}")
    raise RuntimeError("바이낸스 모든 주소 접속 실패")


def upbit(market):
    pts, to = [], None
    for _ in range(60):
        params = {"market": market, "count": 200}
        if to:
            params["to"] = to
        rows = http_get("https://api.upbit.com/v1/candles/days", params=params).json()
        if not rows:
            break
        for c in rows:
            pts.append((parse_date(c["candle_date_time_utc"][:10]), c["trade_price"]))
        oldest = rows[-1]["candle_date_time_utc"]
        to = oldest + "Z" if not oldest.endswith("Z") else oldest
        if len(rows) < 200:
            break
        time.sleep(0.15)
    if not pts:
        raise RuntimeError("업비트 데이터 없음")
    return pts


def gold_monthly():
    """datasets/gold-prices: 1833~1959 Timothy Green, 1960~ World Bank Commodity Markets (월별, USD/oz)."""
    url = "https://raw.githubusercontent.com/datasets/gold-prices/main/data/monthly.csv"
    r = http_get(url)
    lines = r.text.strip().splitlines()
    head = [h.strip().lower() for h in lines[0].split(",")]
    di = head.index("date") if "date" in head else 0
    pi = head.index("price") if "price" in head else 1
    pts = []
    for ln in lines[1:]:
        a = ln.split(",")
        if len(a) > max(di, pi) and a[pi].strip():
            pts.append((parse_date(a[di]), a[pi]))
    return pts


# ---- 엑셀(영국은행 장기통계, Schmelzing 장기금리) 휴리스틱 추출
def download_file(url, suffix):
    r = http_get(url, stream=True, timeout=(10, 120))
    fd, path = tempfile.mkstemp(suffix=suffix)
    with os.fdopen(fd, "wb") as f:
        for chunk in r.iter_content(1 << 20):
            f.write(chunk)
    return path


def excel_candidates(path, header_regex, title_regex=None, header_rows=14):
    """헤더 글자가 header_regex 와 맞는 (시트, 열) 후보 목록."""
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    out = []
    for ws in wb.worksheets:
        try:
            head = list(ws.iter_rows(min_row=1, max_row=header_rows, values_only=True))
        except Exception:  # noqa: BLE001
            continue
        ncol = max((len(r) for r in head), default=0)
        for c in range(ncol):
            txt = " | ".join(str(r[c]).strip() for r in head if c < len(r) and isinstance(r[c], str) and r[c].strip())
            if c == 0 or not txt:
                continue  # 첫 열은 보통 연도/날짜 열
            if re.search(header_regex, txt, re.I) or (title_regex and re.search(title_regex, ws.title, re.I)):
                out.append({"sheet": ws.title, "col": c, "header": txt[:160]})
    return wb, out


def excel_extract(wb, sheet, col, start_row=1):
    """시트에서 (연도/날짜 열, 값 열) 쌍을 읽는다."""
    ws = wb[sheet]
    pts = []
    for row in ws.iter_rows(min_row=start_row, values_only=True):
        if col >= len(row):
            continue
        v = row[col]
        if not isinstance(v, (int, float)) or isinstance(v, bool):
            continue
        d = None
        for c in range(0, min(col, 4)):
            x = row[c]
            if isinstance(x, dt.datetime):
                d = x.date()
                break
            if isinstance(x, dt.date):
                d = x
                break
            if isinstance(x, (int, float)) and 1000 <= x <= 2100 and float(x).is_integer():
                d = dt.date(int(x), 7, 1)
                break
            if isinstance(x, str) and re.match(r"^\d{4}(-\d{2}(-\d{2})?)?$", x.strip()):
                d = parse_date(x.strip())
                break
        if d:
            pts.append((d, v))
    return pts


def pick(cands, *prefs):
    """후보 중 우선순위 정규식에 맞는 것을 고른다."""
    for p in prefs:
        for c in cands:
            if re.search(p, c["sheet"] + " " + c["header"], re.I):
                return c
    return cands[0] if cands else None


BOE_MILLENNIUM = "https://www.bankofengland.co.uk/-/media/boe/files/statistics/research-datasets/a-millennium-of-macroeconomic-data-for-the-uk.xlsx"
SCHMELZING = "https://www.bankofengland.co.uk/-/media/boe/files/working-paper/2020/eight-centuries-of-global-real-interest-rates-r-g-and-the-suprasecular-decline-1311-2018-data.xlsx"
_cache = {}


def boe_millennium_path():
    if "boe" not in _cache:
        _cache["boe"] = download_file(BOE_MILLENNIUM, ".xlsx")
    return _cache["boe"]


def boe_series(header_regex, *prefs, title_regex=None):
    wb, cands = excel_candidates(boe_millennium_path(), header_regex, title_regex)
    log("  후보:", json.dumps(cands[:12], ensure_ascii=False))
    c = pick(cands, *prefs)
    if not c:
        raise RuntimeError("영국은행 장기통계에서 해당 열을 찾지 못함")
    log(f"  선택: 시트={c['sheet']} 열={c['col']} 헤더={c['header']}")
    pts = excel_extract(wb, c["sheet"], c["col"])
    if not pts:
        raise RuntimeError("선택한 열에서 값을 읽지 못함")
    return pts, f"시트 '{c['sheet']}' / 열 '{c['header'][:60]}'"


def boe_bank_rate():
    pts, note = boe_series(r"bank\s*rate", r"bank\s*rate.*(annual|a\d)", r"bank\s*rate", title_regex=r"bank\s*rate")
    last = max(d for d, _ in pts)
    # 최근 구간은 영국은행 통계 DB(IADB)의 일별 공식 기준금리로 이어 붙인다 (실패해도 장기 구간은 유지)
    try:
        r = http_get(
            "https://www.bankofengland.co.uk/boeapps/iadb/fromshowcolumns.asp",
            params={
                "csv.x": "yes", "Datefrom": "01/Jan/1975", "Dateto": "now",
                "SeriesCodes": "IUDBEDR", "CSVF": "TN", "UsingCodes": "Y", "VPD": "Y", "VFD": "N",
            },
        )
        extra = []
        for ln in r.text.strip().splitlines()[1:]:
            a = [x.strip().strip('"') for x in ln.split(",")]
            if len(a) >= 2:
                try:
                    d = dt.datetime.strptime(a[0], "%d %b %Y").date()
                    extra.append((d, float(a[1])))
                except ValueError:
                    continue
        if extra:
            first_daily = min(d for d, _ in extra)
            pts = [p for p in pts if p[0] < first_daily] + extra
            note += " + 영국은행 IADB(IUDBEDR) 일별 1975~"
    except Exception as e:  # noqa: BLE001
        log("  IADB 이어붙이기 실패(장기 구간만 사용):", scrub(e))
        note += f" (최근 구간 IADB 실패, {last.year}년까지)"
    return pts, note


def nl_gov_yield():
    path = download_file(SCHMELZING, ".xlsx")
    wb, cands = excel_candidates(path, r"nether|holland|dutch", title_regex=r"nether|holland|dutch")
    log("  후보:", json.dumps(cands[:12], ensure_ascii=False))
    c = pick(cands, r"nominal", r"nether|holland|dutch")
    if not c:
        raise RuntimeError("Schmelzing 자료에서 네덜란드 열을 찾지 못함")
    log(f"  선택: 시트={c['sheet']} 열={c['col']} 헤더={c['header']}")
    pts = excel_extract(wb, c["sheet"], c["col"])
    if not pts:
        raise RuntimeError("네덜란드 열에서 값을 읽지 못함")
    return pts, f"시트 '{c['sheet']}' / 열 '{c['header'][:60]}'"


# ---------------------------------------------------------------- 지표 정의
def S(id, name, group, unit, fn, *, kind="line", freq="daily", source, url, caution=None, hidden=False, order=0):
    return dict(id=id, name=name, group=group, unit=unit, fn=fn, kind=kind, freq=freq,
                source=source, url=url, caution=caution, hidden=hidden, order=order)


FRED_URL = "https://fred.stlouisfed.org/series/"
SERIES = [
    # 금리
    S("fed_funds", "미국 연방기금금리 (연준)", "rates", "%", lambda: fred("DFF"), kind="step",
      source="FRED(세인트루이스 연준) DFF - Effective Federal Funds Rate, 일별 1954~", url=FRED_URL + "DFF", order=1),
    S("bok_base", "한국은행 기준금리", "rates", "%", ecos_base_rate, kind="step",
      source="한국은행 ECOS 722Y001 / 0101000 (기준금리), 1999.5~", url="https://ecos.bok.or.kr/",
      caution="1999년 5월 이전 자료는 이 출처에 없음", order=2),
    S("boe_bank_rate", "영국은행 기준금리 (Bank Rate)", "rates", "%", boe_bank_rate, kind="step", freq="mixed",
      source="영국은행 'A Millennium of Macroeconomic Data for the UK'(장기, 연 단위) + 영국은행 IADB(IUDBEDR, 일별 1975~)",
      url="https://www.bankofengland.co.uk/statistics/research-datasets",
      caution="엑셀 열 자동 탐색 방식 - 첫 수집 뒤 값을 원본과 대조해 주세요", order=3),
    S("nl_gov_yield", "네덜란드 정부채 금리 (장기)", "rates", "%", nl_gov_yield, kind="line", freq="annual",
      source="Schmelzing(2020), 영국은행 Staff Working Paper No.845 부속 데이터",
      url="https://www.bankofengland.co.uk/working-paper/2020/eight-centuries-of-global-real-interest-rates-r-g-and-the-suprasecular-decline-1311-2018",
      caution="학술 재구성 추정치 / 엑셀 열 자동 탐색 방식 - 원본과 대조 필요", order=4),
    S("us10y", "미국 10년 국채금리", "rates", "%", lambda: fred("DGS10"),
      source="FRED DGS10 (미 재무부), 일별 1962~", url=FRED_URL + "DGS10", order=5),
    S("ecb_rate", "유럽중앙은행 예금금리", "rates", "%", lambda: fred("ECBDFR"), kind="step",
      source="FRED ECBDFR (ECB Deposit Facility Rate), 일별 1999~", url=FRED_URL + "ECBDFR",
      caution="시리즈 코드 확인 필요", order=6),
    S("boj_rate", "일본 단기금리 (콜금리)", "rates", "%", lambda: fred("IRSTCI01JPM156N"), freq="monthly",
      source="FRED IRSTCI01JPM156N (OECD, 일본 콜머니 금리), 월별", url=FRED_URL + "IRSTCI01JPM156N",
      caution="일본은행 정책금리 그 자체가 아님. 시리즈 코드 확인 필요", order=7),
    # 금
    S("gold_usd", "금 가격 (USD/온스, 월별)", "metals", "USD", gold_monthly, freq="monthly",
      source="datasets/gold-prices: 1833~1959 Timothy Green, 1960~ 세계은행 Commodity Markets(월 평균)",
      url="https://github.com/datasets/gold-prices", caution="일별이 아닌 월별 자료", order=1),
    S("gold_gbp", "금 가격 (GBP/온스, 장기)", "metals", "GBP",
      lambda: boe_series(r"gold", r"gold.*price", r"gold", title_regex=r"gold"), freq="annual",
      source="영국은행 'A Millennium of Macroeconomic Data for the UK'",
      url="https://www.bankofengland.co.uk/statistics/research-datasets",
      caution="엑셀 열 자동 탐색 방식 - 단위·시작연도를 원본과 대조해 주세요", order=2),
    S("wti", "WTI 원유 (USD/배럴)", "metals", "USD", lambda: fred("DCOILWTICO"),
      source="FRED DCOILWTICO (미 에너지정보청), 일별 1986~", url=FRED_URL + "DCOILWTICO", order=3),
    # 암호화폐
    S("btc_usd", "비트코인 (USD, 바이낸스)", "crypto", "USD", lambda: binance("BTCUSDT"),
      source="바이낸스 BTCUSDT 일봉 종가(UTC), 2017.8~", url="https://www.binance.com/en/trade/BTC_USDT",
      caution="USDT 기준 (달러와 거의 같지만 정확히 같지는 않음)", order=1),
    S("eth_usd", "이더리움 (USD, 바이낸스)", "crypto", "USD", lambda: binance("ETHUSDT"),
      source="바이낸스 ETHUSDT 일봉 종가(UTC), 2017.8~", url="https://www.binance.com/en/trade/ETH_USDT",
      caution="USDT 기준", order=2),
    S("btc_krw", "비트코인 (원, 업비트)", "crypto", "KRW", lambda: upbit("KRW-BTC"),
      source="업비트 KRW-BTC 일봉 종가(UTC 기준 일봉)", url="https://upbit.com/exchange?code=CRIX.UPBIT.KRW-BTC", order=3),
    S("eth_krw", "이더리움 (원, 업비트)", "crypto", "KRW", lambda: upbit("KRW-ETH"),
      source="업비트 KRW-ETH 일봉 종가(UTC 기준 일봉)", url="https://upbit.com/exchange?code=CRIX.UPBIT.KRW-ETH", order=4),
    # 거시
    S("us_cpi", "미국 소비자물가지수 (CPI)", "macro", "index", lambda: fred("CPIAUCSL"), freq="monthly",
      source="FRED CPIAUCSL (미 노동통계국, 계절조정), 월별 1947~", url=FRED_URL + "CPIAUCSL", order=1),
    # 환율 (1달러당 각국 통화 = USD 기준 환산 후 화면에서 교차환율 계산)
    S("fx_KRW", "원 (1달러당)", "fx", "perUSD", lambda: fred("DEXKOUS"), hidden=True,
      source="FRED DEXKOUS (연준 H.10), 일별 1981~", url=FRED_URL + "DEXKOUS"),
    S("fx_JPY", "엔 (1달러당)", "fx", "perUSD", lambda: fred("DEXJPUS"), hidden=True,
      source="FRED DEXJPUS (연준 H.10), 일별 1971~", url=FRED_URL + "DEXJPUS"),
    S("fx_CNY", "위안 (1달러당)", "fx", "perUSD", lambda: fred("DEXCHUS"), hidden=True,
      source="FRED DEXCHUS (연준 H.10), 일별 1981~", url=FRED_URL + "DEXCHUS"),
    S("fx_EUR", "유로 (1달러당)", "fx", "perUSD", lambda: fred("DEXUSEU", invert=True), hidden=True,
      source="FRED DEXUSEU (연준 H.10, 1유로당 달러)를 역수 변환, 일별 1999~", url=FRED_URL + "DEXUSEU",
      caution="1999년 유로 출범 이전 자료 없음"),
    S("fx_GBP", "파운드 (1달러당)", "fx", "perUSD", lambda: fred("DEXUSUK", invert=True), hidden=True,
      source="FRED DEXUSUK (연준 H.10, 1파운드당 달러)를 역수 변환, 일별 1971~", url=FRED_URL + "DEXUSUK"),
]


def read_old_manifest():
    try:
        with open(os.path.join(DATA_DIR, "manifest.json"), encoding="utf-8") as f:
            return {s["id"]: s for s in json.load(f)["series"]}
    except Exception:  # noqa: BLE001
        return {}


def main():
    only = set(sys.argv[1:])
    os.makedirs(DATA_DIR, exist_ok=True)
    old = read_old_manifest()
    entries, ok_count, fail = [], 0, []
    for s in SERIES:
        if only and s["id"] not in only:
            if s["id"] in old:
                entries.append(old[s["id"]])
            continue
        log(f"[{s['id']}] {s['name']}")
        meta = {k: s[k] for k in ("id", "name", "group", "unit", "kind", "freq", "source", "url", "caution", "hidden", "order")}
        try:
            res = s["fn"]()
            note = None
            if isinstance(res, tuple):
                res, note = res
            pts = clean(res)
            if len(pts) < 2:
                raise RuntimeError("유효한 데이터가 2개 미만")
            path = os.path.join(DATA_DIR, f"{s['id']}.json")
            with open(path, "w", encoding="utf-8") as f:
                json.dump(
                    {"id": s["id"], "t": [day_num(d) for d, _ in pts], "v": [round(v, 6) for _, v in pts]},
                    f, separators=(",", ":"),
                )
            meta.update(status="ok", start=pts[0][0].isoformat(), end=pts[-1][0].isoformat(), n=len(pts),
                        updated=TODAY.isoformat(), detail=note)
            ok_count += 1
            log(f"  OK {len(pts)}개 {pts[0][0]} ~ {pts[-1][0]}")
        except SkipSource as e:
            meta.update(status="skipped", message=scrub(e))
            fail.append((s["id"], f"건너뜀: {scrub(e)}"))
            log("  건너뜀:", scrub(e))
            if s["id"] in old and old[s["id"]].get("n"):
                meta.update({k: old[s["id"]].get(k) for k in ("start", "end", "n", "updated", "detail")})
        except Exception as e:  # noqa: BLE001
            meta.update(status="error", message=scrub(e))
            fail.append((s["id"], scrub(e)))
            log("  실패:", scrub(e))
            log(scrub(traceback.format_exc(limit=3)))
            if s["id"] in old and old[s["id"]].get("n"):  # 이전 데이터 유지
                meta.update({k: old[s["id"]].get(k) for k in ("start", "end", "n", "updated", "detail")})
                meta["status"] = "stale"
        entries.append(meta)

    manifest = {
        "generated": dt.datetime.now(dt.timezone.utc).isoformat(timespec="seconds"),
        "end": END.isoformat(),
        "series": entries,
    }
    with open(os.path.join(DATA_DIR, "manifest.json"), "w", encoding="utf-8") as f:
        json.dump(manifest, f, ensure_ascii=False, indent=1)

    log(f"\n성공 {ok_count} / 전체 {len(entries)}")
    for i, m in fail:
        log(f" - {i}: {m}")
    summary = os.environ.get("GITHUB_STEP_SUMMARY")
    if summary:
        with open(summary, "a", encoding="utf-8") as f:
            f.write(f"### 데이터 수집 결과: 성공 {ok_count} / {len(entries)}\n")
            for e in entries:
                f.write(f"- `{e['id']}` {e.get('status')} {e.get('start','')}~{e.get('end','')} {e.get('message','') or ''}\n")
    if ok_count == 0:
        sys.exit(1)


if __name__ == "__main__":
    main()
